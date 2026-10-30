# Runs the scraper against a tiny local copy of books.toscrape.com
# (html copied from the real site) so the tests don't hit the internet.

import csv
import http.server
import subprocess
import sys
import threading
from functools import partial
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(ROOT))
import scraper

BOOK = '''<li><article class="product_pod">
<div class="image_container"><a href="../../../{slug}/index.html"><img src="../../../../media/cache/{img}" class="thumbnail"></a></div>
<p class="star-rating {rating}"></p>
<h3><a href="../../../{slug}/index.html" title="{title}">{short}</a></h3>
<div class="product_price"><p class="price_color">£{price}</p><p class="instock availability"> In stock </p></div>
</article></li>'''

HIMALAYAS = dict(slug='its-only-the-himalayas_981', img='27/a5/27a53d0bb95bdd88288eaf66c9230d7e.jpg',
                 title="It's Only the Himalayas", short="It's Only the Himalayas", rating='Two', price='45.17')
SHARP = dict(slug='sharp-objects_997', img='32/51/3251cf3a3412f53f339e42cac2134093.jpg',
             title='Sharp Objects', short='Sharp Objects', rating='Four', price='47.82')
DARK = dict(slug='in-a-dark-dark-wood_963', img='23/85/238570a1c284e730dbc737a7e631ae2b.jpg',
            title='In a Dark, Dark Wood', short='In a Dark, Dark ...', rating='One', price='19.63')
STYLES = dict(slug='the-mysterious-affair-at-styles-hercule-poirot-1_452', img='32/94/3294e5eaf73a37958583483fc9a90f04.jpg',
              title='The Mysterious Affair at Styles (Hercule Poirot #1)', short='The Mysterious Affair at ...',
              rating='Four', price='24.80')


def page(books, pager=''):
    items = ''.join(BOOK.format(**b) for b in books)
    return f'<html><body><section><ol class="row">{items}</ol>{pager}</section></body></html>'


SITE = {
    'index.html': '''<html><body><div class="side_categories"><ul><li><a href="catalogue/category/books_1/index.html">Books</a><ul>
        <li><a href="catalogue/category/books/travel_2/index.html"> Travel </a></li>
        <li><a href="catalogue/category/books/mystery_3/index.html"> Mystery </a></li>
        </ul></li></ul></div></body></html>''',
    'catalogue/category/books/travel_2/index.html': page([HIMALAYAS]),
    # mystery has two pages, like on the real site
    'catalogue/category/books/mystery_3/index.html': page(
        [SHARP, DARK], '<ul class="pager"><li class="next"><a href="page-2.html">next</a></li></ul>'),
    'catalogue/category/books/mystery_3/page-2.html': page([STYLES]),
}


@pytest.fixture(scope='module')
def site(tmp_path_factory):
    folder = tmp_path_factory.mktemp('site')
    for path, html in SITE.items():
        (folder / path).parent.mkdir(parents=True, exist_ok=True)
        (folder / path).write_text(html, encoding='utf-8')

    handler = partial(http.server.SimpleHTTPRequestHandler, directory=str(folder))
    handler.log_message = lambda *args: None
    server = http.server.ThreadingHTTPServer(('127.0.0.1', 0), handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    yield f'http://127.0.0.1:{server.server_port}/'
    server.shutdown()


def test_categories(site):
    cats = scraper.get_categories(site)
    assert list(cats) == ['Travel', 'Mystery']


def test_follows_next_page(site):
    url = site + 'catalogue/category/books/mystery_3/index.html'
    books = scraper.scrape_category('Mystery', url, delay=0)
    assert [b['title'] for b in books] == [
        'Sharp Objects',
        'In a Dark, Dark Wood',
        'The Mysterious Affair at Styles (Hercule Poirot #1)',
    ]


def test_book_fields(site):
    url = site + 'catalogue/category/books/mystery_3/index.html'
    book = scraper.scrape_category('Mystery', url, delay=0)[0]
    assert book == {
        'title': 'Sharp Objects',
        'category': 'Mystery',
        'price': 47.82,
        'availability': 'In stock',
        'rating': 4,
        'url': site + 'catalogue/sharp-objects_997/index.html',
        'image': site + 'media/cache/32/51/3251cf3a3412f53f339e42cac2134093.jpg',
    }


def test_full_run(site, tmp_path):
    out = tmp_path / 'output'
    subprocess.run([sys.executable, str(ROOT / 'scraper.py'), '--url', site, '--delay', '0', '--out', str(out)],
                   check=True, capture_output=True)

    with open(out / 'books.csv', encoding='utf-8-sig') as f:
        rows = list(csv.DictReader(f))
    assert len(rows) == 4

    wb = load_workbook(out / 'books.xlsx')
    assert wb['Books'].max_row == 5
    summary = list(wb['Summary'].values)
    assert summary[1] == ('Mystery', 3, 30.75, 19.63, 47.82, 3.0)
    assert summary[2] == ('Travel', 1, 45.17, 45.17, 45.17, 2.0)


def test_single_category(site, tmp_path):
    out = tmp_path / 'output'
    subprocess.run([sys.executable, str(ROOT / 'scraper.py'), '--url', site, '--delay', '0',
                    '--out', str(out), '--category', 'travel'], check=True, capture_output=True)
    with open(out / 'books.csv', encoding='utf-8-sig') as f:
        assert [r['category'] for r in csv.DictReader(f)] == ['Travel']
