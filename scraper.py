import argparse
import csv
import os
import re
import time
from collections import defaultdict
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

BASE_URL = 'https://books.toscrape.com/'
RATINGS = {'One': 1, 'Two': 2, 'Three': 3, 'Four': 4, 'Five': 5}
FIELDS = ['title', 'category', 'price', 'availability', 'rating', 'url', 'image']

session = requests.Session()
session.headers['User-Agent'] = (
    'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 '
    '(KHTML, like Gecko) Chrome/130.0 Safari/537.36'
)


def get_soup(url, retries=3):
    for attempt in range(retries):
        try:
            r = session.get(url, timeout=20)
            r.raise_for_status()
            r.encoding = 'utf-8'  # site doesn't send a charset, so the £ sign breaks without this
            return BeautifulSoup(r.text, 'html.parser')
        except requests.RequestException as e:
            if attempt == retries - 1:
                raise
            print(f'  {e}, retrying...')
            time.sleep(2)


def get_categories(base_url):
    soup = get_soup(base_url)
    links = soup.select('div.side_categories ul li ul li a')
    return {a.text.strip(): urljoin(base_url, a['href']) for a in links}


def parse_book(article, page_url, category):
    link = article.select_one('h3 a')
    price = article.select_one('p.price_color').text
    rating = article.select_one('p.star-rating')['class'][1]

    return {
        # the link text is cut off ("In a Dark, Dark ..."), the title attribute has the full name
        'title': link['title'],
        'category': category,
        'price': float(re.sub(r'[^\d.]', '', price)),
        'availability': article.select_one('p.availability').text.strip(),
        'rating': RATINGS.get(rating, 0),
        'url': urljoin(page_url, link['href']),
        'image': urljoin(page_url, article.select_one('img')['src']),
    }


def scrape_category(name, url, delay=0.5):
    books = []
    while url:
        soup = get_soup(url)
        for article in soup.select('article.product_pod'):
            books.append(parse_book(article, url, name))

        next_page = soup.select_one('li.next a')
        url = urljoin(url, next_page['href']) if next_page else None
        time.sleep(delay)
    return books


def save_csv(books, path):
    with open(path, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.DictWriter(f, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(books)


def save_excel(books, path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Books'
    ws.append(['Title', 'Category', 'Price (GBP)', 'Availability', 'Rating', 'URL', 'Image'])
    for b in books:
        ws.append([b[f] for f in FIELDS])

    for i, width in enumerate([55, 20, 12, 13, 8, 60, 60]):
        ws.column_dimensions['ABCDEFG'[i]].width = width
    for cell in ws['C'][1:]:
        cell.number_format = '£#,##0.00'
    for cell in ws['F'][1:]:
        cell.hyperlink = cell.value
        cell.style = 'Hyperlink'
    style_header(ws)
    ws.auto_filter.ref = ws.dimensions

    # second sheet: quick numbers per category
    summary = wb.create_sheet('Summary')
    summary.append(['Category', 'Books', 'Avg price', 'Min price', 'Max price', 'Avg rating'])
    by_cat = defaultdict(list)
    for b in books:
        by_cat[b['category']].append(b)
    for cat in sorted(by_cat):
        prices = [b['price'] for b in by_cat[cat]]
        ratings = [b['rating'] for b in by_cat[cat]]
        summary.append([cat, len(prices), round(sum(prices) / len(prices), 2),
                        min(prices), max(prices), round(sum(ratings) / len(ratings), 2)])

    for row in summary.iter_rows(min_row=2, min_col=3, max_col=5):
        for cell in row:
            cell.number_format = '£#,##0.00'
    summary.column_dimensions['A'].width = 25
    style_header(summary)

    wb.save(path)


def style_header(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='305496')
    ws.freeze_panes = 'A2'


def main():
    parser = argparse.ArgumentParser(description='Scrape books.toscrape.com to Excel and CSV')
    parser.add_argument('--category', help='only scrape this category, e.g. Travel')
    parser.add_argument('--delay', type=float, default=0.5, help='seconds between requests (default 0.5)')
    parser.add_argument('--out', default='output', help='output folder (default: output)')
    parser.add_argument('--url', default=BASE_URL, help='site url, mainly used by the tests')
    args = parser.parse_args()

    categories = get_categories(args.url)
    if args.category:
        categories = {k: v for k, v in categories.items() if k.lower() == args.category.lower()}
        if not categories:
            raise SystemExit(f'Category "{args.category}" not found')

    books = []
    for name, url in categories.items():
        found = scrape_category(name, url, args.delay)
        print(f'{name}: {len(found)} books')
        books.extend(found)

    os.makedirs(args.out, exist_ok=True)
    save_csv(books, os.path.join(args.out, 'books.csv'))
    save_excel(books, os.path.join(args.out, 'books.xlsx'))
    print(f'\nDone, {len(books)} books saved to {args.out}/')


if __name__ == '__main__':
    main()
